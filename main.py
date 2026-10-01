"""Enterprise Demand Intelligence Platform orchestrator.

Run from the project root with::

    python main.py

The orchestrator owns execution order, run isolation, stage gates, timing,
cache reuse, artifact validation and Power BI integration. Business intelligence
remains in src/*.py and Power BI integration remains in the root-level
powerbi_integration.py module.
"""
from __future__ import annotations

import hashlib
import json
import zipfile
import os
import platform
import shutil
import sys
import threading
import time
import traceback
import uuid
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

import config
import env
import powerbi_integration
from src import analysis, cleaning, forecasting, preprocessing, recommendation_engine

PROJECT_ROOT = Path(__file__).resolve().parent
RUNS_DIRECTORY = PROJECT_ROOT / "data" / "output" / "runs"
PIPELINE_STAGES = (
    "cleaning",
    "preprocessing",
    "analysis",
    "forecasting",
    "recommendation",
    "powerbi",
)

# Cross-run cache is deliberately bounded. A cache hit is safe only when the
# stage cache key, artifact hashes and artifact contract all match.
CACHE_LOOKBACK_RUNS = 40
PARALLEL_ANALYTICS_DEFAULT = True
FILE_HASH_CACHE: dict[tuple[str, int, int], str] = {}
MANIFEST_LOCK = threading.RLock()

STAGE_ARTIFACTS: dict[str, tuple[Path, ...]] = {
    "cleaning": (
        Path("processed") / "processed_dataset.json",
        Path("processed") / "processed_dataset.xlsx",
        Path("processed") / "cleaning_report.json",
    ),
    "preprocessing": (
        Path("output") / "analytical_dataset.json",
        Path("output") / "analytical_dataset.xlsx",
        Path("output") / "preprocessing_report.json",
    ),
    "analysis": (
        Path("output") / "analysis_results.json",
        Path("output") / "analysis_results.xlsx",
        Path("output") / "analysis_report.json",
    ),
    "forecasting": (
        Path("output") / "forecast_results.json",
        Path("output") / "forecast_results.xlsx",
        Path("output") / "forecast_report.json",
    ),
    "recommendation": (
        Path("output") / "recommendation_results.xlsx",
        Path("output") / "recommendation_report.json",
    ),
}

WORKERS: dict[str, Callable[[], Any]] = {
    "cleaning": cleaning.run_cleaning_pipeline,
    "preprocessing": preprocessing.run_preprocessing_pipeline,
    "analysis": analysis.run_analysis,
    "forecasting": forecasting.run_forecasting,
    "recommendation": recommendation_engine.run_recommendation_engine,
}


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def iso_now() -> str:
    return utc_now().isoformat()


def new_run_id() -> str:
    return f"run_{utc_now().strftime('%Y%m%dT%H%M%S_%fZ')}_{uuid.uuid4().hex[:10]}"


def atomic_json(
    path: Path,
    payload: Any,
    *,
    replace_attempts: int = 8,
    replace_backoff_seconds: float = 0.20,
) -> None:
    """Atomically persist JSON with bounded retry for transient Windows locks.

    OneDrive/Windows can briefly hold a destination file during synchronization.
    The retry is deliberately bounded and only applies to the final replace;
    persistent permission failures are still raised.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    temporary.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )

    try:
        for attempt in range(replace_attempts):
            try:
                temporary.replace(path)
                return
            except PermissionError:
                if attempt >= replace_attempts - 1:
                    raise
                time.sleep(replace_backoff_seconds * (attempt + 1))
    finally:
        if temporary.exists():
            try:
                temporary.unlink()
            except OSError:
                pass


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    path = path.resolve()
    stat = path.stat()
    key = (str(path), int(stat.st_size), int(stat.st_mtime_ns))
    cached = FILE_HASH_CACHE.get(key)
    if cached is not None:
        return cached

    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            chunk = handle.read(chunk_size)
            if not chunk:
                break
            digest.update(chunk)
    value = digest.hexdigest()
    FILE_HASH_CACHE[key] = value
    return value


def fingerprint_file(path: Path, include_hash: bool = False) -> dict[str, Any]:
    stat = path.stat()
    result = {
        "path": str(path.resolve()),
        "size": int(stat.st_size),
        "mtime_ns": int(stat.st_mtime_ns),
    }
    if include_hash:
        result["sha256"] = sha256_file(path)
    return result


def fingerprint_paths(paths: list[Path] | tuple[Path, ...]) -> str:
    payload = [
        fingerprint_file(path)
        for path in paths
        if path.is_file()
    ]
    return hashlib.sha256(
        json.dumps(payload, sort_keys=True).encode("utf-8")
    ).hexdigest()


def source_fingerprint() -> str:
    """Fingerprint all executable project code used by the pipeline."""
    paths = [
        PROJECT_ROOT / "main.py",
        PROJECT_ROOT / "env.py",
        PROJECT_ROOT / "config.py",
        PROJECT_ROOT / "powerbi_integration.py",
    ]
    source_dir = PROJECT_ROOT / "src"
    paths.extend(sorted(source_dir.glob("*.py")))

    digest = hashlib.sha256()
    for path in paths:
        if not path.is_file():
            continue
        digest.update(str(path.relative_to(PROJECT_ROOT)).encode("utf-8"))
        digest.update(sha256_file(path).encode("ascii"))
    digest.update(platform.python_version().encode("ascii"))
    return digest.hexdigest()


@dataclass
class StageState:
    name: str
    status: str = "PENDING"
    started_at: str | None = None
    finished_at: str | None = None
    duration_seconds: float = 0.0
    cache_key: str | None = None
    artifacts: list[dict[str, Any]] = field(default_factory=list)
    error: str | None = None
    traceback: str | None = None


@dataclass
class RunContext:
    run_id: str
    dataset_path: Path
    dataset_sha256: str
    run_directory: Path
    processed_directory: Path
    output_directory: Path
    powerbi_root: Path
    powerbi_data_directory: Path
    powerbi_documentation_directory: Path
    manifest_path: Path
    log_path: Path
    source_fingerprint: str
    stages: dict[str, StageState] = field(default_factory=dict)
    log_lines: list[str] = field(default_factory=list)
    powerbi_metadata: dict[str, Any] = field(default_factory=dict)
    pipeline_duration_seconds: float = 0.0

    def to_dict(self) -> dict[str, Any]:
        return {
            "run_id": self.run_id,
            "dataset_path": str(self.dataset_path),
            "dataset_sha256": self.dataset_sha256,
            "run_directory": str(self.run_directory),
            "processed_directory": str(self.processed_directory),
            "output_directory": str(self.output_directory),
            "powerbi_root": str(self.powerbi_root),
            "powerbi_data_directory": str(self.powerbi_data_directory),
            "powerbi_documentation_directory": str(self.powerbi_documentation_directory),
            "source_fingerprint": self.source_fingerprint,
            "stages": {name: vars(state) for name, state in self.stages.items()},
            "powerbi_metadata": self.powerbi_metadata,
            "pipeline_duration_seconds": self.pipeline_duration_seconds,
        }


def log(context: RunContext, message: str) -> None:
    line = f"[{iso_now()}] {message}"
    print(line, flush=True)
    context.log_lines.append(line)
    context.log_path.parent.mkdir(parents=True, exist_ok=True)
    with context.log_path.open("a", encoding="utf-8") as handle:
        handle.write(line + "\n")


def write_manifest(context: RunContext) -> None:
    # Analysis and forecasting may finish on different threads. Manifest writes
    # are serialized so one result cannot overwrite the other.
    with MANIFEST_LOCK:
        atomic_json(context.manifest_path, context.to_dict())


def resolve_dataset() -> Path:
    """Resolve and validate the raw input through the project's loader."""
    from src import data_loader

    path = data_loader.discover_dataset_file().expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Dataset does not exist: {path}")
    if not env.is_supported_dataset_format(path):
        raise ValueError(f"Unsupported dataset format: {path.suffix}")
    if not env.validate_dataset_file_size(path):
        raise ValueError(f"Dataset exceeds the configured size limit: {path}")
    return path


def make_context(dataset: Path) -> RunContext:
    RUNS_DIRECTORY.mkdir(parents=True, exist_ok=True)
    run_id = new_run_id()
    run_directory = RUNS_DIRECTORY / run_id

    # Root Power BI directory is the project's existing, dedicated Power BI
    # workspace. Each run receives unique child directories so successful runs
    # never overwrite each other.
    powerbi_root = Path(getattr(config, "POWERBI_DIRECTORY", PROJECT_ROOT / "powerbi"))
    powerbi_data = Path(getattr(config, "POWERBI_DATA_DIRECTORY", powerbi_root / "data"))
    powerbi_docs = Path(
        getattr(config, "POWERBI_DOCUMENTATION_DIRECTORY", powerbi_root / "documentation")
    )

    context = RunContext(
        run_id=run_id,
        dataset_path=dataset,
        dataset_sha256=sha256_file(dataset),
        run_directory=run_directory,
        processed_directory=run_directory / "processed",
        output_directory=run_directory / "output",
        powerbi_root=powerbi_root,
        powerbi_data_directory=powerbi_data / run_id,
        powerbi_documentation_directory=powerbi_docs / run_id,
        manifest_path=run_directory / "run_manifest.json",
        log_path=run_directory / "execution.log",
        source_fingerprint=source_fingerprint(),
        stages={name: StageState(name) for name in PIPELINE_STAGES},
    )

    for directory in (
        run_directory,
        context.processed_directory,
        context.output_directory,
        context.powerbi_data_directory,
        context.powerbi_documentation_directory,
    ):
        directory.mkdir(parents=True, exist_ok=False)

    # Ensure the project's dedicated Power BI folder exists and is used.
    powerbi_root.mkdir(parents=True, exist_ok=True)
    write_manifest(context)
    return context


def _set_module_attr(
    changes: dict[tuple[Any, str], Any],
    module: Any,
    name: str,
    value: Any,
) -> None:
    if not hasattr(module, name):
        return
    key = (module, name)
    if key not in changes:
        changes[key] = getattr(module, name)
    setattr(module, name, value)


def configure_worker_paths(context: RunContext) -> dict[tuple[Any, str], Any]:
    """Redirect worker output globals for this run and remember every change."""
    changes: dict[tuple[Any, str], Any] = {}

    # The workers use these module-level values for their existing public
    # contracts. We configure them once before any parallel stage begins.
    _set_module_attr(changes, config, "ACTIVE_DATASET_PATH", context.dataset_path)
    _set_module_attr(changes, config, "OUTPUT_DATA_DIRECTORY", context.output_directory)
    _set_module_attr(changes, config, "PROCESSED_DATA_DIRECTORY", context.processed_directory)

    for name, value in {
        "PROCESSED_DIRECTORY": context.processed_directory,
        "PROCESSED_JSON_PATH": context.processed_directory / "processed_dataset.json",
        "PROCESSED_EXCEL_PATH": context.processed_directory / "processed_dataset.xlsx",
        "CLEANING_REPORT_PATH": context.processed_directory / "cleaning_report.json",
    }.items():
        _set_module_attr(changes, cleaning, name, value)

    for name, value in {
        "PROCESSED_DIRECTORY": context.processed_directory,
        "OUTPUT_DIRECTORY": context.output_directory,
        "PROCESSED_JSON_PATH": context.processed_directory / "processed_dataset.json",
        "PROCESSED_EXCEL_PATH": context.processed_directory / "processed_dataset.xlsx",
        "ANALYTICAL_JSON_PATH": context.output_directory / "analytical_dataset.json",
        "ANALYTICAL_EXCEL_PATH": context.output_directory / "analytical_dataset.xlsx",
        "PREPROCESSING_REPORT_PATH": context.output_directory / "preprocessing_report.json",
    }.items():
        _set_module_attr(changes, preprocessing, name, value)

    for name, value in {
        "OUTPUT_DATA_DIRECTORY": context.output_directory,
        "ANALYSIS_WORKBOOK": context.output_directory / "analysis_results.xlsx",
        "ANALYSIS_DATASET_JSON": context.output_directory / "analysis_results.json",
        "ANALYSIS_REPORT_JSON": context.output_directory / "analysis_report.json",
    }.items():
        _set_module_attr(changes, analysis, name, value)

    for name, value in {
        "OUTPUT_DIR": context.output_directory,
        "FORECAST_OUTPUT": context.output_directory / "forecast_results.xlsx",
        "FORECAST_JSON": context.output_directory / "forecast_results.json",
        "FORECAST_REPORT": context.output_directory / "forecast_report.json",
    }.items():
        _set_module_attr(changes, forecasting, name, value)

    for name, value in {
        "OUTPUT_DIR": context.output_directory,
        "ANALYSIS_OUTPUT_FILE": context.output_directory / "analysis_results.xlsx",
        "FORECAST_OUTPUT_FILE": context.output_directory / "forecast_results.xlsx",
        "RECOMMENDATION_OUTPUT_FILE": context.output_directory / "recommendation_results.xlsx",
        "RECOMMENDATION_REPORT_FILE": context.output_directory / "recommendation_report.json",
    }.items():
        _set_module_attr(changes, recommendation_engine, name, value)

    return changes


def restore_worker_paths(changes: dict[tuple[Any, str], Any]) -> None:
    for (module, name), old_value in reversed(list(changes.items())):
        setattr(module, name, old_value)


def expected_artifacts(context: RunContext, stage: str) -> list[Path]:
    roots = {
        "processed": context.processed_directory,
        "output": context.output_directory,
    }
    return [
        roots[item.parts[0]] / Path(*item.parts[1:])
        for item in STAGE_ARTIFACTS[stage]
    ]


def artifact_state(
    paths: list[Path],
    include_hash: bool = True,
    run_id: str | None = None,
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for path in paths:
        if not path.is_file():
            continue
        record = fingerprint_file(path, include_hash=include_hash)
        if run_id is not None:
            record["run_id"] = run_id
        records.append(record)
    return records


def _path_is_within(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def _validate_artifact_content(path: Path) -> tuple[bool, str | None]:
    """Validate the minimum structural contract without assuming business schema."""
    if not path.is_file():
        return False, "file does not exist"
    if path.stat().st_size <= 0:
        return False, "file is empty"

    suffix = path.suffix.lower()
    try:
        if suffix == ".json":
            with path.open("r", encoding="utf-8") as handle:
                payload = json.load(handle)
            if payload is None:
                return False, "JSON payload is null"
        elif suffix in {".xlsx", ".xlsm"}:
            # openpyxl performs ZIP/container and workbook-structure checks.
            from openpyxl import load_workbook
            if not zipfile.is_zipfile(path):
                return False, "Excel file is not a valid ZIP workbook container"
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                if not workbook.sheetnames:
                    return False, "Excel workbook contains no worksheets"
            finally:
                workbook.close()
        else:
            return True, None
    except Exception as exc:
        return False, f"content validation failed: {type(exc).__name__}: {exc}"
    return True, None


def validate_stage_artifacts(
    context: RunContext,
    stage: str,
    paths: list[Path] | None = None,
    *,
    require_current_run: bool = True,
) -> tuple[bool, list[dict[str, Any]], str | None]:
    """Validate a stage's complete artifact contract and ownership boundary."""
    expected = paths if paths is not None else expected_artifacts(context, stage)
    if len(expected) != len(STAGE_ARTIFACTS.get(stage, ())):
        return False, [], f"{stage} artifact contract count mismatch."

    records: list[dict[str, Any]] = []
    for path in expected:
        if require_current_run and not _path_is_within(path, context.run_directory):
            return False, [], f"{stage} artifact is outside the current run directory: {path}"
        valid, reason = _validate_artifact_content(path)
        if not valid:
            return False, [], f"{stage} artifact invalid: {path} ({reason})"
        record = fingerprint_file(path, include_hash=True)
        record["run_id"] = context.run_id
        records.append(record)

    return True, records, None


def validate_stage_dependencies(context: RunContext, stage: str) -> tuple[bool, str | None]:
    """Fail closed unless every required upstream artifact is from this run."""
    dependency_stages = {
        "cleaning": (),
        "preprocessing": ("cleaning",),
        "analysis": ("preprocessing",),
        "forecasting": ("preprocessing",),
        "recommendation": ("analysis", "forecasting"),
        "powerbi": ("cleaning", "preprocessing", "analysis", "forecasting", "recommendation"),
    }
    for upstream in dependency_stages.get(stage, ()):
        upstream_state = context.stages[upstream]
        if upstream_state.status not in {"PASSED", "REUSED"}:
            return False, f"required upstream stage {upstream} is not complete: {upstream_state.status}"
        valid, _records, error = validate_stage_artifacts(context, upstream)
        if not valid:
            return False, f"required upstream stage {upstream} failed artifact validation: {error}"
    return True, None


def stage_dependency_paths(context: RunContext, stage: str) -> list[Path]:
    if stage == "cleaning":
        return [context.dataset_path]
    if stage == "preprocessing":
        return [context.processed_directory / "processed_dataset.json"]
    if stage in {"analysis", "forecasting"}:
        # JSON is the canonical computational artifact. Excel remains the
        # presentation/Power BI artifact and is not a required cache dependency.
        return [context.output_directory / "analytical_dataset.json"]
    if stage == "recommendation":
        return [
            context.output_directory / "analysis_results.xlsx",
            context.output_directory / "forecast_results.xlsx",
        ]
    return []


def stage_cache_key(context: RunContext, stage: str) -> str:
    dependencies = stage_dependency_paths(context, stage)
    payload = {
        "stage": stage,
        "dataset_sha256": context.dataset_sha256,
        "source_fingerprint": context.source_fingerprint,
        "dependencies": [
            fingerprint_file(path)
            for path in dependencies
            if path.is_file()
        ],
        "artifact_contract": [str(path) for path in STAGE_ARTIFACTS.get(stage, ())],
    }
    return hashlib.sha256(
        json.dumps(payload, sort_keys=True, default=str).encode("utf-8")
    ).hexdigest()


def validate_artifact_records(
    records: list[dict[str, Any]],
    owner_run_directory: Path | None = None,
    owner_run_id: str | None = None,
) -> bool:
    if not records:
        return False
    for record in records:
        path = Path(record.get("path", ""))
        if owner_run_directory is not None and not _path_is_within(path, owner_run_directory):
            return False
        if owner_run_id is not None and record.get("run_id") != owner_run_id:
            return False
        if not path.is_file() or path.stat().st_size <= 0:
            return False
        if int(path.stat().st_size) != int(record.get("size", -1)):
            return False
        expected_hash = record.get("sha256")
        if not expected_hash or sha256_file(path) != expected_hash:
            return False
        valid, _ = _validate_artifact_content(path)
        if not valid:
            return False
    return True


def cached_stage_is_valid(run_directory: Path, stage: str, cache_key: str) -> bool:
    manifest_path = run_directory / "run_manifest.json"
    if not manifest_path.is_file():
        return False
    try:
        payload = json.loads(manifest_path.read_text(encoding="utf-8"))
        state = payload.get("stages", {}).get(stage, {})
        if state.get("status") not in {"PASSED", "REUSED"}:
            return False
        if state.get("cache_key") != cache_key:
            return False
        records = state.get("artifacts", [])
        if len(records) != len(STAGE_ARTIFACTS[stage]):
            return False
        return validate_artifact_records(
            records,
            owner_run_directory=run_directory,
            owner_run_id=payload.get("run_id", run_directory.name),
        )
    except Exception:
        return False


def find_cached_stage(context: RunContext, stage: str, cache_key: str) -> Path | None:
    if not RUNS_DIRECTORY.is_dir():
        return None
    candidates = sorted(
        (
            path
            for path in RUNS_DIRECTORY.iterdir()
            if path.is_dir() and path.name != context.run_id
        ),
        key=lambda path: path.stat().st_mtime_ns,
        reverse=True,
    )
    for candidate in candidates[:CACHE_LOOKBACK_RUNS]:
        if cached_stage_is_valid(candidate, stage, cache_key):
            return candidate
    return None


def reuse_cached_stage(context: RunContext, stage: str, cached_run: Path) -> None:
    current = expected_artifacts(context, stage)
    manifest = json.loads((cached_run / "run_manifest.json").read_text(encoding="utf-8"))
    records = manifest["stages"][stage]["artifacts"]
    if len(records) != len(current):
        raise RuntimeError(f"Cached {stage} artifact contract mismatch.")
    cached_run_id = manifest.get("run_id", cached_run.name)
    if not validate_artifact_records(
        records, owner_run_directory=cached_run, owner_run_id=cached_run_id
    ):
        raise RuntimeError(f"Cached {stage} artifacts failed integrity/ownership validation.")

    for record, destination in zip(records, current):
        source = Path(record["path"])
        destination.parent.mkdir(parents=True, exist_ok=True)
        # Copy instead of hard-link: each run owns independent artifact files.
        shutil.copy2(source, destination)


def _mark_stage_failed(
    context: RunContext,
    stage: str,
    started: float,
    exc: BaseException,
    *,
    persist_manifest: bool = True,
) -> None:
    state = context.stages[stage]
    state.status = "FAILED"
    state.error = f"{type(exc).__name__}: {exc}"
    state.traceback = traceback.format_exc()
    state.duration_seconds = time.perf_counter() - started
    state.finished_at = iso_now()
    log(context, f"{stage.upper()}: FAILED in {state.duration_seconds:.2f}s — {state.error}")
    if persist_manifest:
        write_manifest(context)


def run_worker_stage(
    context: RunContext,
    stage: str,
    *,
    persist_manifest: bool = True,
) -> bool:
    state = context.stages[stage]
    state.status = "RUNNING"
    state.started_at = iso_now()
    state.cache_key = stage_cache_key(context, stage)
    state.error = None
    state.traceback = None
    state.artifacts = []
    if persist_manifest:
        write_manifest(context)

    started = time.perf_counter()
    try:
        dependencies_ok, dependency_error = validate_stage_dependencies(context, stage)
        if not dependencies_ok:
            raise RuntimeError(dependency_error or f"{stage} upstream artifact contract failed")

        cached = find_cached_stage(context, stage, state.cache_key)
        if cached is not None:
            reuse_cached_stage(context, stage, cached)
            valid, records, error = validate_stage_artifacts(context, stage)
            if not valid:
                raise RuntimeError(f"Reused {stage} artifacts failed current-run validation: {error}")
            state.status = "REUSED"
            state.artifacts = records
            state.duration_seconds = time.perf_counter() - started
            state.finished_at = iso_now()
            log(context, f"{stage.upper()}: REUSED from {cached.name} in {state.duration_seconds:.2f}s")
            if persist_manifest:
                write_manifest(context)
            return True

        log(context, f"{stage.upper()}: START")
        WORKERS[stage]()
        paths = expected_artifacts(context, stage)
        valid, records, error = validate_stage_artifacts(context, stage)
        if not valid:
            raise RuntimeError(error or f"{stage} completed with invalid artifacts")

        state.status = "PASSED"
        state.artifacts = records
        state.duration_seconds = time.perf_counter() - started
        state.finished_at = iso_now()
        log(context, f"{stage.upper()}: PASSED in {state.duration_seconds:.2f}s")
        if persist_manifest:
            write_manifest(context)
        return True
    except Exception as exc:
        _mark_stage_failed(
            context,
            stage,
            started,
            exc,
            persist_manifest=persist_manifest,
        )
        return False


def run_analysis_and_forecasting(context: RunContext) -> bool:
    """Run independent read-only analytical stages concurrently when enabled."""
    parallel_value = os.getenv(
        "EDI_PARALLEL_ANALYTICS",
        "1" if PARALLEL_ANALYTICS_DEFAULT else "0",
    ).strip().lower()
    parallel = parallel_value not in {"0", "false", "no", "off"}

    if not parallel:
        return (
            run_worker_stage(context, "analysis")
            and run_worker_stage(context, "forecasting")
        )

    log(context, "ANALYTICS: running ANALYSIS and FORECASTING in parallel")
    with ThreadPoolExecutor(
        max_workers=2,
        thread_name_prefix="edi-analytics",
    ) as pool:
        futures = {
            stage: pool.submit(
                run_worker_stage,
                context,
                stage,
                persist_manifest=False,
            )
            for stage in ("analysis", "forecasting")
        }

        # Workers deliberately do not persist the shared run manifest while
        # executing concurrently. The orchestrator collects both outcomes and
        # performs one serialized manifest commit on the main thread.
        analysis_ok = futures["analysis"].result()
        forecasting_ok = futures["forecasting"].result()

    write_manifest(context)
    return analysis_ok and forecasting_ok


def run_powerbi_stage(context: RunContext) -> bool:
    state = context.stages["powerbi"]
    state.status = "RUNNING"
    state.started_at = iso_now()
    state.error = None
    state.traceback = None
    write_manifest(context)

    started = time.perf_counter()
    try:
        dependencies_ok, dependency_error = validate_stage_dependencies(context, "powerbi")
        if not dependencies_ok:
            raise RuntimeError(dependency_error or "Power BI upstream artifact contract failed")
        log(context, "POWERBI: START")
        ok, metadata, error = powerbi_integration.generate_powerbi(
            context,
            lambda message: log(context, message),
        )
        if not ok:
            raise RuntimeError(error or "Power BI generation failed")

        project = Path(metadata["project"])
        integration_manifest = project / "integration_manifest.json"
        if not project.is_dir():
            raise RuntimeError(f"Power BI project directory was not created: {project}")
        if not integration_manifest.is_file():
            raise RuntimeError("Power BI integration manifest was not created.")
        if not Path(metadata.get("pbip", "")).is_file():
            raise RuntimeError("Power BI .pbip project descriptor was not created.")
        if not Path(metadata.get("report", {}).get("definition", "")).is_dir():
            raise RuntimeError("Power BI report definition directory was not created.")

        context.powerbi_metadata = metadata
        state.status = "PASSED"
        state.artifacts = [
            {
                "path": str(integration_manifest.resolve()),
                "size": integration_manifest.stat().st_size,
                "mtime_ns": integration_manifest.stat().st_mtime_ns,
                "sha256": sha256_file(integration_manifest),
            }
        ]
        state.duration_seconds = time.perf_counter() - started
        state.finished_at = iso_now()
        write_manifest(context)
        log(context, f"POWERBI: PASSED in {state.duration_seconds:.2f}s")

        if os.getenv("EDI_OPEN_POWERBI", "0").strip().lower() in {"1", "true", "yes", "on"}:
            powerbi_integration.maybe_open_powerbi_project(
                context,
                project,
                lambda message: log(context, message),
            )
        return True
    except Exception as exc:
        _mark_stage_failed(context, "powerbi", started, exc)
        return False


def print_timing_summary(context: RunContext, total_seconds: float) -> None:
    print("\n" + "=" * 78)
    print("PIPELINE EXECUTION TIME SUMMARY")
    print("=" * 78)
    print(f"{'Stage':<18}{'Status':<12}{'Time (s)':>12}{'Time':>16}")
    print("-" * 78)
    for name in PIPELINE_STAGES:
        state = context.stages[name]
        seconds = state.duration_seconds
        print(
            f"{name.upper():<18}"
            f"{state.status:<12}"
            f"{seconds:>12.2f}"
            f"{format_duration(seconds):>16}"
        )
    print("-" * 78)

    stage_sum = sum(context.stages[name].duration_seconds for name in PIPELINE_STAGES)
    analytics_path = max(
        context.stages["analysis"].duration_seconds,
        context.stages["forecasting"].duration_seconds,
    )
    critical_path = (
        context.stages["cleaning"].duration_seconds
        + context.stages["preprocessing"].duration_seconds
        + analytics_path
        + context.stages["recommendation"].duration_seconds
        + context.stages["powerbi"].duration_seconds
    )

    print(f"{'STAGE-TIME SUM':<46}{stage_sum:>12.2f}s")
    print(f"{'PARALLEL CRITICAL PATH':<46}{critical_path:>12.2f}s")
    print(f"{'TOTAL WALL-CLOCK':<46}{total_seconds:>12.2f}s")
    if stage_sum > 0:
        print(f"{'PARALLELIZATION SAVING':<46}{max(0.0, stage_sum - critical_path):>12.2f}s")
        print(f"{'ACTUAL WALL-CLOCK SAVING':<46}{max(0.0, stage_sum - total_seconds):>12.2f}s")
    print("=" * 78)


def format_duration(seconds: float) -> str:
    seconds = max(0.0, float(seconds))
    if seconds < 60:
        return f"{seconds:.1f}s"
    minutes, remainder = divmod(seconds, 60)
    return f"{int(minutes)}m {remainder:04.1f}s"


def run_application() -> int:
    pipeline_started = time.perf_counter()
    context: RunContext | None = None
    changes: dict[tuple[Any, str], Any] = {}

    if not env.authenticate():
        return 1

    env.ensure_project_directories()
    dataset = resolve_dataset()
    context = make_context(dataset)

    try:
        log(context, f"RUN {context.run_id} START")
        log(context, f"Dataset: {dataset}")
        log(context, f"Dataset SHA-256: {context.dataset_sha256}")
        log(context, f"Power BI workspace: {context.powerbi_root}")

        changes = configure_worker_paths(context)

        if not run_worker_stage(context, "cleaning"):
            return 1
        if not run_worker_stage(context, "preprocessing"):
            return 1
        if not run_analysis_and_forecasting(context):
            return 1
        if not run_worker_stage(context, "recommendation"):
            return 1
        if not run_powerbi_stage(context):
            return 1

        log(context, "PIPELINE: SUCCESS")
        return 0
    finally:
        total_seconds = time.perf_counter() - pipeline_started
        if context is not None:
            context.pipeline_duration_seconds = total_seconds
            restore_worker_paths(changes)
            write_manifest(context)
            print_timing_summary(context, total_seconds)


def main() -> int:
    try:
        return run_application()
    except KeyboardInterrupt:
        print("\nPipeline interrupted by user.")
        return 130
    except Exception as exc:
        print(
            f"Fatal pipeline error: {type(exc).__name__}: {exc}",
            file=sys.stderr,
        )
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
